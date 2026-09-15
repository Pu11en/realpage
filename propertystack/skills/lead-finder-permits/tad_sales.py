"""T5: a county appraisal-district "improved sales" zip -- the same
county-appraisal-zip shape as tad_zip.py's permits file, but this one
covers sold buildings with a sale price instead of new-construction
permits, and the apartment sheet's name changes between years ("Apartment"
in the 2026 file, "Apartments" in the 2025 file) so the match is a
case-insensitive prefix instead of an exact worksheet name.

The sheet has no buyer/grantee column and no city column at all (checked
live against both the 2025 and 2026 files) -- only a deed document number.
So `buyer` always stays blank here (never guessed), and `city` falls back
to the recipe's own `county` name, same fallback find_sold.py uses when a
county sales file doesn't carry a per-row city.
"""
from __future__ import annotations

import datetime
import io
import sys
import zipfile
from pathlib import Path
from typing import Callable

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "lead-finder"))
from merge import merge_records  # noqa: E402
from record import LeadRecord  # noqa: E402

FetchBytes = Callable[[str], bytes]

_EXCEL_EPOCH = datetime.date(1899, 12, 30)


def default_fetch_bytes(url: str) -> bytes:
    import urllib.request

    req = urllib.request.Request(url, headers={"User-Agent": "propertystack-live-self-test"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return resp.read()


def fetch_sale_rows(recipe: dict, fetch_bytes: FetchBytes = default_fetch_bytes) -> list[dict]:
    """Download every zip URL in recipe["zip_urls"], unzip the single xlsx
    inside each, and return every row (as a plain dict keyed by the sheet's
    own header row) from whichever worksheet's name starts with
    recipe["sheet_name_prefix"] (case-insensitive, default "apartment")."""
    import openpyxl

    sheet_prefix = recipe.get("sheet_name_prefix", "apartment").lower()
    rows: list[dict] = []
    for url in recipe.get("zip_urls", []):
        try:
            data = fetch_bytes(url)
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                xlsx_names = [n for n in zf.namelist() if n.lower().endswith((".xlsx", ".xlsm"))]
                if not xlsx_names:
                    continue
                xlsx_bytes = zf.read(xlsx_names[0])
        except Exception:
            continue
        try:
            wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
        except Exception:
            continue
        sheet_name = next((n for n in wb.sheetnames if n.lower().startswith(sheet_prefix)), None)
        if sheet_name is None:
            continue
        ws = wb[sheet_name]
        header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(header, excel_row))
            row["_source_url"] = url
            rows.append(row)
    return rows


def find_apartment_sales(
    area: str,
    recipe: dict,
    fetch_bytes: FetchBytes = default_fetch_bytes,
    today: datetime.date | None = None,
) -> list[LeadRecord]:
    """One LeadRecord (stage="sold") per apartment-sheet row with a known
    unit count >= min_units and a sale date on or after recipe["since"]."""
    today = today or datetime.date.today()
    units_field = recipe.get("units_field", "Number of Units")
    date_field = recipe.get("date_field", "Document Date")
    name_field = recipe.get("name_field", "Site Name")
    address_field = recipe.get("address_field", "Address")
    price_field = recipe.get("price_field", "Adjusted Sale Price")
    fallback_price_field = recipe.get("fallback_price_field", "Contract Sale Price")
    doc_field = recipe.get("doc_field", "Document Number")
    min_units = recipe.get("min_units", 20)
    since = _to_date(recipe.get("since", "2024-09-01")) or datetime.date(2024, 9, 1)
    county = recipe.get("county", "")

    records = []
    for row in fetch_sale_rows(recipe, fetch_bytes):
        units = _to_int(row.get(units_field))
        if units is None or units < min_units:
            continue
        sale_date = _to_date(row.get(date_field))
        if sale_date is None or not (since <= sale_date <= today):
            continue
        price = _to_int(row.get(price_field)) or _to_int(row.get(fallback_price_field))
        name = str(row.get(name_field) or "").strip()
        address = str(row.get(address_field) or "").strip()
        source_url = row.get("_source_url", "")
        doc_number = str(row.get(doc_field) or "").strip()
        why = f"recently sold apartment property ({units} units"
        why += f", ${price:,})" if price else ", price not public)"
        records.append(
            LeadRecord(
                name=name,
                area=area,
                city=county,
                address=address,
                units=units,
                stage="sold",
                sale_date=sale_date.isoformat(),
                links={"sales_file": source_url, "deed_doc": doc_number},
                sources=[{"fact": "sale_date", "url": source_url}],
                why=why,
            )
        )
    return merge_records(records)


def _to_int(value) -> int | None:
    try:
        result = int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None
    return result if result > 0 else None


def _to_date(value) -> datetime.date | None:
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    if text.replace(".", "", 1).isdigit():
        try:
            return _EXCEL_EPOCH + datetime.timedelta(days=int(float(text)))
        except (ValueError, OverflowError):
            return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None
