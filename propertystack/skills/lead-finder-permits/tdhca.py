"""T5: TDHCA's two statewide affordable-housing spreadsheets.

- HTC Property Inventory (plain xlsx, not zipped): the "PropInventory" sheet
  holds every tax-credit property TDHCA has ever funded; new-construction
  awards are found by filtering ConType == "New Construction" and the
  "Year" column (the award year -- confirmed live against real rows this
  is the field that tracks the award, not the oddly-small "Board Approval"
  column that also exists on this sheet) >= recipe["since_year"]. No phone
  field on this sheet.
- 4% (non-competitive) HTC status log (plain xlsx, one workbook per update):
  the real header row isn't row 1 -- the file leads with a title block and
  footnotes, so the header is found by scanning for the row whose first
  cell is literally "TDHCA Number" (confirmed live: row 11 in the
  2026-08-03 file) rather than assuming a fixed row number. Filtered to
  Construction Type == "NC"; carries "Applicant Phone" (the plan's
  affordable-pipeline phone) and Total Units.

Both are plain xlsx downloads (no zip), unlike every other T3/T4/T5 county
file -- `fetch_bytes` here returns the workbook bytes directly.
"""
from __future__ import annotations

import datetime
import io
import sys
from pathlib import Path
from typing import Callable

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "lead-finder"))
from merge import merge_records  # noqa: E402
from record import LeadRecord  # noqa: E402

FetchBytes = Callable[[str], bytes]


def default_fetch_bytes(url: str) -> bytes:
    import urllib.request

    req = urllib.request.Request(url, headers={"User-Agent": "propertystack-live-self-test"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return resp.read()


def _load_workbook(url: str, fetch_bytes: FetchBytes):
    import openpyxl

    try:
        data = fetch_bytes(url)
    except Exception:
        return None
    try:
        return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        return None


def _find_htc_inventory_records(area: str, recipe: dict, fetch_bytes: FetchBytes) -> list[LeadRecord]:
    url = recipe.get("htc_inventory_url")
    if not url:
        return []
    wb = _load_workbook(url, fetch_bytes)
    if wb is None:
        return []
    sheet_name = recipe.get("htc_inventory_sheet", "PropInventory")
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    since_year = int(recipe.get("since_year", 2024))
    records = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, excel_row))
        if str(row.get("ConType") or "").strip() != "New Construction":
            continue
        year = _to_int(row.get("Year"))
        if year is None or year < since_year:
            continue
        units = _to_int(row.get("Total Units"))
        name = str(row.get("Development Name") or "").strip()
        city = str(row.get("Project City") or "").strip()
        address = str(row.get("Project Address") or "").strip()
        records.append(
            LeadRecord(
                name=name,
                area=area,
                city=city,
                address=address,
                units=units,
                stage="planned",
                links={"htc_inventory": url},
                sources=[{"fact": "units", "url": url}],
                why=f"tax-credit new-construction award ({year})",
            )
        )
    return records


def _find_status_log_records(area: str, recipe: dict, fetch_bytes: FetchBytes) -> list[LeadRecord]:
    records = []
    for url in recipe.get("status_log_urls", []):
        wb = _load_workbook(url, fetch_bytes)
        if wb is None:
            continue
        ws = wb[wb.sheetnames[0]]
        header = None
        header_row_idx = None
        for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
            if row and str(row[0] or "").strip() == "TDHCA Number":
                header = [str(v or "").strip() for v in row]
                header_row_idx = idx
                break
        if header is None:
            continue
        for excel_row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            row = dict(zip(header, excel_row))
            if str(row.get("Construction Type") or "").strip() != "NC":
                continue
            units = _to_int(row.get("Total Units"))
            name = str(row.get("Development Name") or "").strip()
            city = str(row.get("Development City") or "").strip()
            address = str(row.get("Development Address") or "").strip()
            phone = str(row.get("Applicant Phone") or "").strip()
            records.append(
                LeadRecord(
                    name=name,
                    area=area,
                    city=city,
                    address=address,
                    units=units,
                    stage="planned",
                    office_phone=phone,
                    links={"status_log": url},
                    sources=[{"fact": "units", "url": url}],
                    why="4% (non-competitive) HTC affordable pipeline application",
                )
            )
    return records


def find_new_affordable_projects(
    area: str,
    recipe: dict,
    fetch_bytes: FetchBytes = default_fetch_bytes,
    today: datetime.date | None = None,
) -> list[LeadRecord]:
    """Combine TDHCA's HTC inventory (new-construction tax-credit awards
    since recipe["since_year"]) and the 4% status log (in-progress
    affordable applications, with applicant phone) into one deduped list."""
    records = _find_htc_inventory_records(area, recipe, fetch_bytes)
    records += _find_status_log_records(area, recipe, fetch_bytes)
    return merge_records(records)


def _to_int(value) -> int | None:
    try:
        result = int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None
    return result if result > 0 else None
